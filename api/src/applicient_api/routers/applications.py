"""F6/F7 — application CRUD, the pipeline board's state machine, and
the application-agent trigger/resume streams. Same
EventSourceResponse/SSE stage-event shape as radar.py/job_groups.py
for the two streaming endpoints; everything else is plain REST.
"""

from __future__ import annotations

import asyncio
import contextlib
import io
import json
import uuid
from datetime import datetime, timezone

import httpx
import websockets
from fastapi import APIRouter, Depends, HTTPException, Response, WebSocket, WebSocketDisconnect
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func
from sqlalchemy.orm import Session, sessionmaker
from sse_starlette.sse import EventSourceResponse

from applicient_agents.application_service import (
    BROWSER_WORKER_URL,
    request_cancel,
    resume_application_attempt,
    start_application_attempt,
)

from applicient_api import pipeline_stage_service, schemas
from applicient_api.deps import current_user_id, get_db, get_session_factory
from applicient_api.document_resolution import available_document_types, resolve_application_documents
from applicient_api.models.agents import AgentRun, RunEvent
from applicient_api.models.discovery import Job
from applicient_api.models.enums import EventActor
from applicient_api.models.pipeline import Application, ApplicationAttempt, ApplicationEvent
from applicient_api.models.profile import Persona
from applicient_api.models.scoring import FitScore
from applicient_api.object_storage import get_object
from applicient_api.credit_ledger import FEATURE_APPLICATION_APPLY, require_credits
from applicient_api.pipeline_service import MarkAppliedError, TransitionError, is_ghosted, latest_event_times, transition
from applicient_api.rate_limit import rate_limit
from applicient_api.pipeline_service import mark_applied as mark_applied_service

# An attempt in any of these statuses is actively holding a real
# browser-worker session or awaiting a human decision — deleting the
# application out from under it would orphan that session with nothing
# left to ever resolve it. Blocks the delete instead (Adrian's own new
# "delete from pipeline" ask still needs this guard, same discipline
# `cancel_attempt`'s own two-branch shape already applies elsewhere).
_ACTIVE_ATTEMPT_STATUSES = ("in_progress", "awaiting_review", "awaiting_handoff", "awaiting_email")

router = APIRouter(prefix="/applications", tags=["applications"])


def _to_out(application: Application, *, ghosted: bool, job: Job | None = None) -> schemas.ApplicationOut:
    out = schemas.ApplicationOut.model_validate(application)
    out.ghosted = ghosted
    if job is not None:
        out.job_title = job.title
        out.company_name = job.company_name_raw
    return out


def _jobs_by_id(db: Session, applications: list[Application]) -> dict[uuid.UUID, Job]:
    job_ids = {a.job_id for a in applications}
    if not job_ids:
        return {}
    return {j.id: j for j in db.query(Job).filter(Job.id.in_(job_ids)).all()}


def _latest_fit_scores(
    db: Session, applications: list[Application]
) -> dict[tuple[uuid.UUID, uuid.UUID], FitScore]:
    """Keyed by (job_id, persona_id) — FitScore is append-only history
    (rescored whenever the profile/persona changes), so this is the
    latest row per pair, not the only one. Used by the export below to
    show each application's real fit score/recommendation, not just
    its pipeline stage."""

    pairs = {(a.job_id, a.persona_id) for a in applications}
    if not pairs:
        return {}
    job_ids = {p[0] for p in pairs}
    persona_ids = {p[1] for p in pairs}
    rows = (
        db.query(FitScore)
        .filter(FitScore.job_id.in_(job_ids), FitScore.persona_id.in_(persona_ids))
        .order_by(FitScore.created_at.desc())
        .all()
    )
    latest: dict[tuple[uuid.UUID, uuid.UUID], FitScore] = {}
    for fs in rows:
        key = (fs.job_id, fs.persona_id)
        if key in pairs and key not in latest:
            latest[key] = fs
    return latest


def _build_applications_workbook(
    applications: list[Application],
    jobs: dict[uuid.UUID, Job],
    fit_scores: dict[tuple[uuid.UUID, uuid.UUID], FitScore],
    latest_events: dict[uuid.UUID, datetime],
) -> Response:
    """Adrian, direct: "I dont want it to be csv. I want it to be
    spreadsheet instead. The tables created must be formatted nicely
    easy to read. And the exported table should like have the info
    like the salary, etc." — real openpyxl formatting (a colored bold
    header row, frozen header + autofilter, zebra striping, real
    numeric types for salary/score so a spreadsheet's own SUM/AVERAGE
    work on them) rather than a bare CSV dump, plus every column Job
    Inbox itself shows, not just the pipeline-specific fields the old
    export had."""

    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"

    headers = [
        "Job title", "Company", "Location", "Remote policy", "Employment type",
        "Salary min", "Salary max", "Salary currency", "Recommendation", "Fit score",
        "Stage", "Autonomy level", "Ghosted", "Applied at", "Discovered at", "Apply URL",
    ]
    ws.append(headers)
    header_fill = PatternFill(start_color="0F62FE", end_color="0F62FE", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    for a in applications:
        job = jobs.get(a.job_id)
        fs = fit_scores.get((a.job_id, a.persona_id))
        ghosted = is_ghosted(a, latest_event_at=latest_events.get(a.id))
        ws.append([
            job.title if job else "",
            job.company_name_raw if job else "",
            job.location if job else "",
            job.remote_policy if job else "",
            job.employment_type if job else "",
            float(job.salary_min) if job and job.salary_min is not None else None,
            float(job.salary_max) if job and job.salary_max is not None else None,
            job.salary_currency if job else "",
            fs.recommendation if fs else "",
            float(fs.overall_score) if fs else None,
            a.state,
            a.autonomy_level or "",
            "Yes" if ghosted else "No",
            a.applied_at.replace(tzinfo=None) if a.applied_at else None,
            a.created_at.replace(tzinfo=None),
            job.apply_url if job else "",
        ])

    # Real date formatting (not the raw ISO-ish default openpyxl would
    # otherwise show) for the two datetime columns.
    for row in ws.iter_rows(min_row=2, min_col=14, max_col=15):
        for cell in row:
            if cell.value is not None:
                cell.number_format = "yyyy-mm-dd hh:mm"

    # Zebra striping — real readability, not just a header/body split.
    stripe_fill = PatternFill(start_color="F4F4F4", end_color="F4F4F4", fill_type="solid")
    for row_idx in range(2, ws.max_row + 1):
        if row_idx % 2 == 0:
            for cell in ws[row_idx]:
                cell.fill = stripe_fill

    for col_idx, header in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(header)
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, min_row=2):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 40)

    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=applications.xlsx"},
    )


@router.post("", response_model=schemas.ApplicationOut, status_code=201)
def create_application(
    body: schemas.ApplicationCreate,
    db: Session = Depends(get_db),
    user_id: uuid.UUID = Depends(current_user_id),
) -> schemas.ApplicationOut:
    job = db.query(Job).filter_by(id=body.job_id, user_id=user_id).one_or_none()
    if job is None:
        raise HTTPException(404, "job not found")

    existing = db.query(Application).filter_by(job_id=body.job_id, user_id=user_id).one_or_none()
    if existing is not None:
        return _to_out(existing, ghosted=False, job=job)

    persona = db.query(Persona).filter_by(id=body.persona_id, user_id=user_id).one_or_none()
    if persona is None:
        raise HTTPException(404, "persona not found")

    application = Application(
        user_id=user_id,
        job_id=body.job_id,
        persona_id=persona.id,
        job_group_id=body.job_group_id,
        primary_document_id=body.primary_document_id,
        # M5 follow-up — the user's own first pipeline stage, not a
        # fixed "discovered" literal. The fallback is defensive only:
        # every user should already have stages by the time they can
        # reach this endpoint (provisioned at signup/seed).
        state=pipeline_stage_service.first_stage_key(db, user_id=user_id) or "discovered",
    )
    db.add(application)
    db.commit()
    db.refresh(application)
    db.add(
        ApplicationEvent(
            application_id=application.id,
            actor=EventActor.USER.value,
            event_type="created",
            payload={},
            occurred_at=datetime.now(timezone.utc),
        )
    )
    db.commit()
    return _to_out(application, ghosted=False, job=job)


@router.get("", response_model=list[schemas.ApplicationOut])
def list_applications(
    state: str | None = None,
    db: Session = Depends(get_db),
    user_id: uuid.UUID = Depends(current_user_id),
) -> list[schemas.ApplicationOut]:
    query = db.query(Application).filter_by(user_id=user_id)
    if state:
        query = query.filter_by(state=state)
    applications = query.order_by(Application.created_at.desc()).all()
    latest = latest_event_times(db, [a.id for a in applications])
    jobs = _jobs_by_id(db, applications)
    return [
        _to_out(a, ghosted=is_ghosted(a, latest_event_at=latest.get(a.id)), job=jobs.get(a.job_id))
        for a in applications
    ]


@router.get("/export")
def export_applications(
    format: str = "xlsx",
    db: Session = Depends(get_db),
    user_id: uuid.UUID = Depends(current_user_id),
) -> Response:
    """F7.6. `format=json` stays for API/debug parity; the real export
    surface is `xlsx` now (Adrian, direct — "I dont want it to be
    csv") — see _build_applications_workbook's own docstring."""

    applications = db.query(Application).filter_by(user_id=user_id).order_by(Application.created_at.desc()).all()
    latest = latest_event_times(db, [a.id for a in applications])
    jobs = _jobs_by_id(db, applications)

    if format == "json":
        fit_scores = _latest_fit_scores(db, applications)
        rows = [
            {
                "id": str(a.id),
                "job_id": str(a.job_id),
                "job_title": jobs[a.job_id].title if a.job_id in jobs else "",
                "company_name": jobs[a.job_id].company_name_raw if a.job_id in jobs else "",
                "location": jobs[a.job_id].location if a.job_id in jobs else None,
                "salary_min": float(jobs[a.job_id].salary_min) if a.job_id in jobs and jobs[a.job_id].salary_min is not None else None,
                "salary_max": float(jobs[a.job_id].salary_max) if a.job_id in jobs and jobs[a.job_id].salary_max is not None else None,
                "salary_currency": jobs[a.job_id].salary_currency if a.job_id in jobs else None,
                "recommendation": fit_scores[(a.job_id, a.persona_id)].recommendation
                if (a.job_id, a.persona_id) in fit_scores
                else None,
                "state": a.state,
                "ghosted": is_ghosted(a, latest_event_at=latest.get(a.id)),
                "autonomy_level": a.autonomy_level,
                "applied_at": a.applied_at.isoformat() if a.applied_at else None,
                "created_at": a.created_at.isoformat(),
            }
            for a in applications
        ]
        return Response(content=json.dumps(rows, indent=2), media_type="application/json")

    fit_scores = _latest_fit_scores(db, applications)
    return _build_applications_workbook(applications, jobs, fit_scores, latest)


@router.get("/{application_id}", response_model=schemas.ApplicationDetailOut)
def get_application(
    application_id: uuid.UUID,
    db: Session = Depends(get_db),
    user_id: uuid.UUID = Depends(current_user_id),
) -> schemas.ApplicationDetailOut:
    application = db.query(Application).filter_by(id=application_id, user_id=user_id).one_or_none()
    if application is None:
        raise HTTPException(404, "application not found")
    events = (
        db.query(ApplicationEvent)
        .filter_by(application_id=application_id)
        .order_by(ApplicationEvent.occurred_at.desc())
        .all()
    )
    attempts = (
        db.query(ApplicationAttempt)
        .filter_by(application_id=application_id)
        .order_by(ApplicationAttempt.attempt_number.desc())
        .all()
    )
    latest = latest_event_times(db, [application_id])
    ghosted = is_ghosted(application, latest_event_at=latest.get(application_id))
    job = db.get(Job, application.job_id)
    out = schemas.ApplicationDetailOut.model_validate(application)
    out.ghosted = ghosted
    if job is not None:
        out.job_title = job.title
        out.company_name = job.company_name_raw
    out.events = [schemas.ApplicationEventOut.model_validate(e) for e in events]
    out.attempts = [schemas.ApplicationAttemptOut.model_validate(a) for a in attempts]
    out.available_documents = available_document_types(db, application=application)
    return out


@router.delete("/{application_id}", status_code=204)
def delete_application(
    application_id: uuid.UUID,
    db: Session = Depends(get_db),
    user_id: uuid.UUID = Depends(current_user_id),
):
    """Adrian, direct: "I want to also be able to delete from
    pipeline." Blocked while an attempt is actively holding a live
    browser-worker session or a pending human decision — see
    _ACTIVE_ATTEMPT_STATUSES above — so this can't silently orphan a
    running agent. `ApplicationEvent`/`ApplicationAttempt` cascade-
    delete at the DB level (ondelete="CASCADE", models/pipeline.py);
    CalendarEvent/InterviewSession's own application_id just goes
    null (ondelete="SET NULL"), never deleted."""

    application = db.query(Application).filter_by(id=application_id, user_id=user_id).one_or_none()
    if application is None:
        raise HTTPException(404, "application not found")
    active_attempt = (
        db.query(ApplicationAttempt.id)
        .filter(
            ApplicationAttempt.application_id == application_id,
            ApplicationAttempt.status.in_(_ACTIVE_ATTEMPT_STATUSES),
        )
        .first()
    )
    if active_attempt is not None:
        raise HTTPException(409, "cancel the running agent before deleting this application")
    db.delete(application)
    db.commit()


@router.get("/{application_id}/documents/{doc_type}")
def get_application_document(
    application_id: uuid.UUID,
    doc_type: str,
    user_id: uuid.UUID = Depends(current_user_id),
    session_factory: sessionmaker = Depends(get_session_factory),
) -> Response:
    """Same fallback chain the application-agent's own `browser_upload`
    tool uses (`document_resolution.resolve_application_documents`): an
    explicit primary document, then the job_group's tailored documents,
    then the persona's own originally-uploaded CV. Backs the Pipeline
    page's "apply by email" download buttons, which used to require a
    job_group_id (a Composer session behind this application) and hit
    the identical dead end an Inbox-direct application gave the agent —
    nothing to offer either path. Ownership is enforced inside the
    resolver itself (it filters the Application by user_id first), not
    duplicated here — an unowned or nonexistent application_id just
    resolves to nothing, same 404 as a genuinely missing document."""

    resolved = resolve_application_documents(session_factory, application_id=application_id, user_id=user_id)
    doc = resolved.get(doc_type)
    if doc is None:
        raise HTTPException(404, f"no '{doc_type}' document available for this application")
    return Response(
        content=get_object(doc.key),
        media_type=doc.content_type,
        headers={"Content-Disposition": f'attachment; filename="{doc.filename}"'},
    )


@router.patch("/{application_id}", response_model=schemas.ApplicationOut)
def update_application(
    application_id: uuid.UUID,
    body: schemas.ApplicationUpdate,
    db: Session = Depends(get_db),
    user_id: uuid.UUID = Depends(current_user_id),
) -> schemas.ApplicationOut:
    application = db.query(Application).filter_by(id=application_id, user_id=user_id).one_or_none()
    if application is None:
        raise HTTPException(404, "application not found")

    if body.state is not None:
        try:
            transition(db, application=application, new_state=body.state, actor=EventActor.USER.value)
        except TransitionError as exc:
            raise HTTPException(400, str(exc))
    if body.autonomy_level is not None:
        application.autonomy_level = body.autonomy_level
    if body.primary_document_id is not None:
        application.primary_document_id = body.primary_document_id
    db.commit()
    db.refresh(application)
    latest = latest_event_times(db, [application_id])
    job = db.get(Job, application.job_id)
    return _to_out(application, ghosted=is_ghosted(application, latest_event_at=latest.get(application_id)), job=job)


@router.post("/{application_id}/mark-applied", response_model=schemas.ApplicationOut)
def mark_applied(
    application_id: uuid.UUID,
    body: schemas.MarkAppliedIn,
    db: Session = Depends(get_db),
    user_id: uuid.UUID = Depends(current_user_id),
) -> schemas.ApplicationOut:
    """F9.3 — manual mark-as-applied from any job card, independent of
    whether the agent ever ran for this application at all."""

    application = db.query(Application).filter_by(id=application_id, user_id=user_id).one_or_none()
    if application is None:
        raise HTTPException(404, "application not found")

    try:
        mark_applied_service(
            db,
            application=application,
            actor=EventActor.USER.value,
            event_type="marked_applied_manually",
            note=body.note or "",
        )
    except MarkAppliedError as exc:
        raise HTTPException(409, str(exc))
    db.commit()
    db.refresh(application)
    return _to_out(application, ghosted=False, job=db.get(Job, application.job_id))


@router.post(
    "/{application_id}/apply",
    dependencies=[
        Depends(rate_limit("application-apply", limit=10, window_seconds=60)),
        Depends(require_credits(FEATURE_APPLICATION_APPLY, label="applying to this job")),
    ],
)
def apply(
    application_id: uuid.UUID,
    db: Session = Depends(get_db),
    user_id: uuid.UUID = Depends(current_user_id),
) -> EventSourceResponse:
    """Starts a real application-agent run. Streams
    `{"event": "stage" | "interrupt" | "done" | "error", "data": "<json>"}`
    — an `interrupt` event means the stream has paused for a human
    decision; resolve it via `POST /applications/{id}/attempts/{attempt_id}/resume`."""

    application = db.query(Application).filter_by(id=application_id, user_id=user_id).one_or_none()
    if application is None:
        raise HTTPException(404, "application not found")
    return EventSourceResponse(
        start_application_attempt(get_session_factory(), application_id=application_id, user_id=user_id)
    )


@router.post("/{application_id}/attempts/{attempt_id}/resume")
def resume(
    application_id: uuid.UUID,
    attempt_id: uuid.UUID,
    body: schemas.InterruptDecisionsIn,
    user_id: uuid.UUID = Depends(current_user_id),
) -> EventSourceResponse:
    decisions = [
        {"type": d.type, **({"message": d.message} if d.message is not None else {})} for d in body.decisions
    ]
    return EventSourceResponse(
        resume_application_attempt(
            get_session_factory(), attempt_id=attempt_id, user_id=user_id, decisions=decisions
        )
    )


@router.post("/{application_id}/attempts/{attempt_id}/cancel")
def cancel_attempt(
    application_id: uuid.UUID,
    attempt_id: uuid.UUID,
    db: Session = Depends(get_db),
    user_id: uuid.UUID = Depends(current_user_id),
):
    """Stop button — cooperative: a real Task.cancel() on whatever the
    application-agent is doing right now (an LLM call or a live
    browser-worker tool call) if a live run is still holding this
    attempt (the common case, via application_service.request_cancel),
    or a direct DB mark if not (the stream already disconnected on its
    own — same two-branch discipline as streaming.py's generic
    /agent-runs/{id}/cancel). Closing the browser-worker session itself
    only happens on the live path; the fallback has no http client
    left to do it with, same gap the existing disconnect-cleanup path
    already has."""

    attempt = _owned_attempt(db, application_id=application_id, attempt_id=attempt_id, user_id=user_id)
    if attempt.status != "in_progress":
        raise HTTPException(409, f"attempt is already {attempt.status}, nothing to cancel")

    if not request_cancel(attempt_id):
        attempt.status = "cancelled"
        attempt.finished_at = datetime.now(timezone.utc)
        if attempt.agent_run_id is not None:
            run = db.get(AgentRun, attempt.agent_run_id)
            if run is not None and run.status == "running":
                run.status = "cancelled"
                run.finished_at = datetime.now(timezone.utc)
                next_seq = (db.query(func.max(RunEvent.seq)).filter_by(agent_run_id=run.id).scalar() or 0) + 1
                db.add(
                    RunEvent(
                        user_id=user_id, agent_run_id=run.id, seq=next_seq,
                        event_type="cancelled", data={"attempt_id": str(attempt_id)},
                    )
                )
        db.commit()
        db.refresh(attempt)
    return {"status": attempt.status}


def _owned_attempt(
    db: Session, *, application_id: uuid.UUID, attempt_id: uuid.UUID, user_id: uuid.UUID
) -> ApplicationAttempt:
    attempt = (
        db.query(ApplicationAttempt)
        .join(Application, ApplicationAttempt.application_id == Application.id)
        .filter(
            ApplicationAttempt.id == attempt_id,
            ApplicationAttempt.application_id == application_id,
            Application.user_id == user_id,
        )
        .one_or_none()
    )
    if attempt is None:
        raise HTTPException(404, "attempt not found")
    return attempt


@router.get("/{application_id}/attempts/{attempt_id}/events", response_model=schemas.RunEventsOut)
def list_attempt_events(
    application_id: uuid.UUID,
    attempt_id: uuid.UUID,
    since_seq: int = 0,
    db: Session = Depends(get_db),
    user_id: uuid.UUID = Depends(current_user_id),
):
    """Replay/reconnect endpoint — mirrors saved_searches.py's
    `list_run_events` exactly, same reason: navigating away from an
    in-progress run and back used to show nothing but `status:
    "in_progress"`, because the live log only ever existed on the one
    open SSE connection (application_service.py's `_emit` now writes
    every event down as a durable `RunEvent` row first). Polling this
    on an interval while the attempt is still active approximates a
    live reconnect — same "no separate live-push channel" tradeoff as
    the radar run-events endpoint."""

    attempt = _owned_attempt(db, application_id=application_id, attempt_id=attempt_id, user_id=user_id)

    run_status = "completed"
    if attempt.agent_run_id is not None:
        run = db.get(AgentRun, attempt.agent_run_id)
        if run is not None:
            run_status = run.status

    events = (
        db.query(RunEvent)
        .filter(RunEvent.agent_run_id == attempt.agent_run_id, RunEvent.seq > since_seq)
        .order_by(RunEvent.seq)
        .all()
    )
    return schemas.RunEventsOut(
        run_status=run_status,
        events=[schemas.RunEventOut.model_validate(e) for e in events],
    )


@router.get("/{application_id}/attempts/{attempt_id}/screenshots/{index}")
def get_attempt_screenshot(
    application_id: uuid.UUID,
    attempt_id: uuid.UUID,
    index: int,
    db: Session = Depends(get_db),
    user_id: uuid.UUID = Depends(current_user_id),
) -> Response:
    """F6.7's review screen needs the actual screenshot, not just a
    key string — `index` supports Python-style negative indexing
    (`-1` = latest) so the Pipeline panel can always ask for "the most
    recent one" without first fetching the attempt to count them."""

    attempt = _owned_attempt(db, application_id=application_id, attempt_id=attempt_id, user_id=user_id)
    keys = attempt.screenshot_keys or []
    if not keys or not (-len(keys) <= index < len(keys)):
        raise HTTPException(404, "screenshot not found")
    return Response(content=get_object(keys[index]), media_type="image/png")


@router.get("/{application_id}/attempts/{attempt_id}/live-snapshot")
async def get_attempt_live_snapshot(
    application_id: uuid.UUID,
    attempt_id: uuid.UUID,
    session_id: str,
    db: Session = Depends(get_db),
    user_id: uuid.UUID = Depends(current_user_id),
) -> dict:
    """F6.7's review screen, field-by-field: while a `submit_application`
    (or any other) interrupt is pending, the browser session named in
    its own args is still open and paused — this proxies browser-worker's
    live accessibility-tree snapshot for it, so the human reviewing the
    approve/reject decision sees the actual current state of every
    field, not a reconstruction. Proxied (not called directly from the
    browser) so browser-worker — which has no auth of its own — stays
    unreachable from outside the compose network."""

    _owned_attempt(db, application_id=application_id, attempt_id=attempt_id, user_id=user_id)
    async with httpx.AsyncClient(timeout=15) as client:
        try:
            r = await client.get(f"{BROWSER_WORKER_URL}/sessions/{session_id}/snapshot")
        except httpx.HTTPError as exc:
            raise HTTPException(502, f"browser-worker unreachable: {exc}")
    if r.is_error:
        raise HTTPException(r.status_code, r.text)
    return r.json()


@router.websocket("/{application_id}/attempts/{attempt_id}/live-browser")
async def live_browser_proxy(
    websocket: WebSocket,
    application_id: uuid.UUID,
    attempt_id: uuid.UUID,
    session_id: str,
    db: Session = Depends(get_db),
    user_id: uuid.UUID = Depends(current_user_id),
) -> None:
    """The dedicated Live Browser page (§8.3/F6.5) — proxies
    browser-worker's own screencast WebSocket
    (`browser_worker/main.py`'s `/sessions/{id}/screencast`) through
    `api` rather than exposing browser-worker's internal port directly
    to the frontend: it has no auth of its own and isn't meant to be
    reachable outside the compose network, same reasoning as the
    live-snapshot proxy above. A dumb pipe in both directions — real
    PNG frames pass straight through one way, JSON input-forwarding
    events (click/type/key) pass straight through the other — all the
    actual logic lives in browser-worker's own SessionManager."""

    try:
        _owned_attempt(db, application_id=application_id, attempt_id=attempt_id, user_id=user_id)
    except HTTPException:
        await websocket.close(code=4404)
        return

    await websocket.accept()
    worker_ws_url = BROWSER_WORKER_URL.replace("http://", "ws://").replace("https://", "wss://")
    try:
        async with websockets.connect(f"{worker_ws_url}/sessions/{session_id}/screencast") as upstream:

            async def pump_frames() -> None:
                async for frame in upstream:
                    if isinstance(frame, (bytes, bytearray)):
                        await websocket.send_bytes(frame)

            frames_task = asyncio.create_task(pump_frames())
            try:
                while True:
                    message = await websocket.receive_text()
                    await upstream.send(message)
            except WebSocketDisconnect:
                pass
            finally:
                frames_task.cancel()
                with contextlib.suppress(asyncio.CancelledError):
                    await frames_task
    except (OSError, websockets.exceptions.WebSocketException):
        with contextlib.suppress(RuntimeError):
            await websocket.close(code=1011)
